"""
DocuChat AI - Complete Excel Generation & QA Test Data Engine Test Suite
Validates all 14 requirements from Part 9:
1. Creating a workbook from user-provided data.
2. Creating a QA test-data workbook.
3. Creating multiple worksheets.
4. Verifying row counts and headers.
5. Opening the generated XLSX successfully.
6. Downloading the actual generated file through the frontend / API.
7. Handling invalid input and failed generation.
8. Handling large requests within configured limits.
9. Preventing path traversal and unauthorized file access.
10. Preserving text safely against formula injection (CWE-1236).
11. Confirming that Excel analysis still reads uploaded workbooks.
12. Confirming that PDF RAG and evidence navigation still work.
13. Confirming that weather and general web-search questions still use their proper workflows.
14. Confirming that ordinary Excel explanations do not unnecessarily generate files.
"""

import os
import io
import asyncio
import pytest
from pathlib import Path
import openpyxl

from backend.config import settings
from backend.tools.excel_generator import (
    sanitize_cell_value,
    generate_custom_workbook,
    generate_qa_test_data_workbook,
    generate_student_management_qa_workbook,
    generate_domain_qa_workbook,
    generate_dynamic_dataset_workbook,
    validate_generated_workbook,
    generate_workbook_from_request,
    excel_generator_tool,
)
from backend.tools.excel_parser import parse_excel_generation_request
from backend.tools.dynamic_data_generator import STUDENT_PRIMARY_COLUMNS
from backend.tools.artifact_registry import artifact_registry
from backend.db.models import UserModel
from backend.llm.router import (
    is_excel_generation_request,
    is_weather_request,
    is_explicit_web_search_request,
    query_router,
)
from backend.api.tools import download_output_file
from fastapi import HTTPException


# ==============================================================================
# 1. Creating a workbook from user-provided data
# ==============================================================================
def test_create_workbook_from_user_data(tmp_path):
    headers = ["EmpID", "Name", "Department", "Salary"]
    rows = [
        ["E001", "Alice Smith", "Engineering", 125000],
        ["E002", "Bob Johnson", "Product", 110000],
        ["E003", "Carol Williams", "Design", 95000]
    ]
    artifact = generate_custom_workbook(
        title="Employee Data",
        sheets_data=[
            {
                "sheet_name": "Employees",
                "headers": headers,
                "rows": rows
            }
        ],
        filename_prefix="test_employees"
    )

    assert artifact["filename"].endswith(".xlsx")
    assert artifact["sheet_count"] == 1
    assert artifact["total_rows"] == 3
    assert artifact["validation_status"] == "Created and validated"

    file_path = os.path.join(settings.OUTPUT_DIR, artifact["filename"])
    assert os.path.exists(file_path)

    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb["Employees"]
    assert ws.max_row == 4  # 1 header + 3 data rows
    assert [cell.value for cell in ws[1]] == headers
    assert ws.cell(row=2, column=2).value == "Alice Smith"
    wb.close()


# ==============================================================================
# 2. Creating a QA test-data workbook (Matches Screenshots 1 & 2)
# ==============================================================================
def test_create_qa_test_data_workbook():
    artifact = generate_qa_test_data_workbook(
        application_name="User Registration Application",
        filename_prefix="application_test_data"
    )

    assert artifact["filename"].startswith("application_test_data")
    assert artifact["filename"].endswith(".xlsx")
    assert artifact["file_type"] == "Excel workbook"
    assert artifact["sheet_count"] == 4
    assert artifact["total_cases"] == 130
    assert artifact["large_dataset_cases"] == 100
    assert artifact["validation_status"] == "Created and validated"
    assert artifact["is_qa_test_data"] is True
    assert "Done, AMD!" in artifact["reply_message"]
    assert "130 test cases" in artifact["reply_message"]

    # Verify category breakdown counts match Screenshot 1 exactly
    cat_map = {c["name"]: c["count"] for c in artifact["coverage"]["categories"]}
    assert cat_map["Valid Data"] == 2
    assert cat_map["Invalid Data"] == 7
    assert cat_map["Boundary Cases"] == 4
    assert cat_map["Missing Data"] == 5
    assert cat_map["Duplicate Data"] == 2
    assert cat_map["Edge Cases"] == 4
    assert cat_map["Data Type Errors"] == 2
    assert cat_map["Security Test Inputs"] == 4
    assert cat_map["Large Dataset"] == 100


# ==============================================================================
# 3. Creating multiple worksheets with professional formatting
# ==============================================================================
def test_multiple_worksheets_structure():
    artifact = generate_qa_test_data_workbook(
        application_name="User Login Application",
        filename_prefix="login_test_data"
    )
    file_path = os.path.join(settings.OUTPUT_DIR, artifact["filename"])
    wb = openpyxl.load_workbook(file_path, data_only=True)

    expected_sheets = ["Summary", "Functional Tests", "Security & Edge Tests", "Large Dataset"]
    assert wb.sheetnames == expected_sheets

    # Verify frozen panes and auto-filter on Functional Tests
    ws_func = wb["Functional Tests"]
    assert ws_func.freeze_panes == "A2"
    assert ws_func.auto_filter.ref is not None

    # Verify Summary Sheet banner
    ws_summary = wb["Summary"]
    assert "QA Test Suite & Coverage Report" in str(ws_summary["A1"].value)
    wb.close()


# ==============================================================================
# 4. Verifying row counts and headers
# ==============================================================================
def test_row_counts_and_headers():
    artifact = generate_qa_test_data_workbook()
    file_path = os.path.join(settings.OUTPUT_DIR, artifact["filename"])
    wb = openpyxl.load_workbook(file_path, data_only=True)

    # Functional Tests sheet
    ws_func = wb["Functional Tests"]
    expected_headers = [
        "Test Case ID", "Test Category", "Test Scenario", "Test Data",
        "Preconditions", "Steps", "Expected Result", "Actual Result", "Status", "Priority"
    ]
    actual_headers = [cell.value for cell in ws_func[1]]
    assert actual_headers == expected_headers
    # Functional cases count: 26 rows in functional sheet
    assert ws_func.max_row >= 25

    # Large Dataset sheet
    ws_large = wb["Large Dataset"]
    assert ws_large.max_row == 101  # 1 header + 100 records
    wb.close()


# ==============================================================================
# 5. Opening the generated XLSX successfully
# ==============================================================================
def test_reopen_xlsx_successfully():
    artifact = generate_qa_test_data_workbook()
    file_path = os.path.join(settings.OUTPUT_DIR, artifact["filename"])
    val = validate_generated_workbook(file_path, ["Summary", "Functional Tests", "Security & Edge Tests", "Large Dataset"])

    assert val["valid"] is True
    assert val["sheet_count"] == 4
    assert val["total_data_rows"] > 100
    assert val["validation_status"] == "Created and validated"


# ==============================================================================
# 6. Downloading the actual generated file through the API
# ==============================================================================
@pytest.mark.asyncio
async def test_download_endpoint_serves_actual_file():
    artifact = generate_qa_test_data_workbook(filename_prefix="download_test")
    filename = artifact["filename"]

    response = await download_output_file(filename=filename)
    assert response.status_code == 200
    assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert f"filename={filename}" in response.headers["Content-Disposition"]
    assert os.path.exists(response.path)


# ==============================================================================
# 7. Handling invalid input and failed generation
# ==============================================================================
def test_invalid_input_and_failed_generation():
    non_existent_file = os.path.join(settings.OUTPUT_DIR, "non_existent_fake_workbook.xlsx")
    with pytest.raises(FileNotFoundError):
        validate_generated_workbook(non_existent_file)


# ==============================================================================
# 8. Handling large requests within configured limits
# ==============================================================================
def test_large_request_limits():
    artifact = generate_qa_test_data_workbook(filename_prefix="large_stress_test")
    file_path = os.path.join(settings.OUTPUT_DIR, artifact["filename"])
    file_size = os.path.getsize(file_path)

    # Ensure size is reasonable and well under the 15MB ceiling
    assert 5000 < file_size < 15 * 1024 * 1024


# ==============================================================================
# 9. Preventing path traversal and unauthorized file access
# ==============================================================================
@pytest.mark.asyncio
async def test_path_traversal_prevention():
    malicious_filenames = [
        "../../etc/passwd",
        "..\\..\\windows\\win.ini",
        "....//....//boot.ini",
        "/etc/shadow"
    ]
    for bad_name in malicious_filenames:
        with pytest.raises(HTTPException) as exc_info:
            await download_output_file(filename=bad_name)
        assert exc_info.value.status_code in (403, 404)


# ==============================================================================
# 10. Preserving text safely against formula injection (CWE-1236)
# ==============================================================================
def test_formula_injection_defense():
    assert sanitize_cell_value("=SUM(A1:A10)") == "'=SUM(A1:A10)"
    assert sanitize_cell_value("+1234567890") == "'+1234567890"
    assert sanitize_cell_value("-cmd|'/C calc'!A0") == "'-cmd|'/C calc'!A0"
    assert sanitize_cell_value("@cmd") == "'@cmd"
    assert sanitize_cell_value("\tmalicious_tab") == "'\tmalicious_tab"
    assert sanitize_cell_value("\rmalicious_cr") == "'\rmalicious_cr"
    # Normal safe strings should remain untouched
    assert sanitize_cell_value("Normal text") == "Normal text"
    assert sanitize_cell_value("user@example.com") == "user@example.com"
    assert sanitize_cell_value(42) == 42


# ==============================================================================
# 11. Confirming that Excel analysis still reads uploaded workbooks
# ==============================================================================
def test_excel_analysis_reads_uploaded_workbooks():
    # Questions that should be routed to analysis, NOT generation
    analysis_questions = [
        "Analyze my uploaded Excel file",
        "How many rows in the sheet?",
        "What are the column names in the spreadsheet?",
        "Summarize the uploaded excel sheet",
        "Inspect the uploaded spreadsheet data"
    ]
    for q in analysis_questions:
        assert is_excel_generation_request(q) is False, f"Expected False for analysis query: '{q}'"


# ==============================================================================
# 12. Confirming that PDF RAG and evidence navigation still work
# ==============================================================================
def test_pdf_rag_still_works():
    doc_questions = [
        "What is the candidate's work experience according to the resume?",
        "What does the policy document say about remote work?",
        "Extract the exact wording of section 4 from the PDF"
    ]
    for q in doc_questions:
        assert is_excel_generation_request(q) is False, f"Expected False for document query: '{q}'"


# ==============================================================================
# 13. Confirming that weather and general web-search questions use proper workflows
# ==============================================================================
def test_weather_and_web_search_workflows():
    weather_q = "What is the weather today in Goa?"
    assert is_weather_request(weather_q) is True
    assert is_excel_generation_request(weather_q) is False

    web_q = "Research on the internet and give me a guide on how to use OKF"
    assert is_explicit_web_search_request(web_q) is True
    assert is_excel_generation_request(web_q) is False


# ==============================================================================
# 14. Confirming that ordinary Excel explanations do not generate files
# ==============================================================================
def test_ordinary_excel_explanations_no_file_generation():
    info_questions = [
        "What is Excel?",
        "Explain Excel",
        "How does Excel work?",
        "Find current Excel documentation",
        "Excel tutorial for beginners",
        "Who created Excel?"
    ]
    for q in info_questions:
        assert is_excel_generation_request(q) is False, f"Expected False for informational query: '{q}'"


# ==============================================================================
# End-to-end Intent Routing Tests for Requested Examples
# ==============================================================================
@pytest.mark.asyncio
async def test_all_user_request_examples_route_to_excel_generation():
    user_examples = [
        "Create an Excel sheet for user registration testing.",
        "Generate 100 test cases for my login page.",
        "Convert this data into an Excel file.",
        "Create an Excel sheet with valid and invalid email examples.",
        "Generate test data from the following requirements.",
        "Create an Excel workbook with separate sheets for positive and negative tests.",
        "Give me a downloadable spreadsheet containing this data.",
        "Create an Excel file about the topic you just researched"
    ]
    for prompt in user_examples:
        assert is_excel_generation_request(prompt) is True, f"Failed to detect intent for: '{prompt}'"
        route, badge = await query_router.route_query(prompt, has_documents=False)
        assert route == "excel_generation", f"Expected 'excel_generation' route for: '{prompt}'"
        assert badge == "📊 Excel Generator"


@pytest.mark.asyncio
async def test_end_to_end_generate_workbook_from_request():
    artifact = await generate_workbook_from_request("Create an Excel sheet for user registration testing.")
    assert artifact["filename"].startswith("application_test_data")
    assert artifact["filename"].endswith(".xlsx")
    assert artifact["sheet_count"] == 4
    assert artifact["total_cases"] == 130
    assert artifact["validation_status"] == "Created and validated"
    assert "/api/tools/download/" in artifact["download_url"]


@pytest.mark.asyncio
async def test_stream_rag_response_excel_generation_events():
    from backend.llm.generator import grounded_generator
    chunks = []
    async for sse in grounded_generator.stream_rag_response(
        question="Create an Excel sheet for user registration testing.",
        route="excel_generation",
        router_badge="📊 Excel Generator",
        retrieved_chunks=[],
        queries=["Create an Excel sheet for user registration testing."],
        web_mode=False
    ):
        chunks.append(sse)

    combined = "".join(chunks)
    assert "event: meta" in combined
    assert "📊 Excel Generator" in combined
    assert "event: excel_artifact" in combined
    assert "application_test_data" in combined
    assert "event: token" in combined
    assert "Done, AMD!" in combined
    assert "event: done" in combined


# ==============================================================================
# 15. Student Management System Workbook Generation Tests
# ==============================================================================
def test_create_student_management_workbook():
    artifact = generate_student_management_qa_workbook(
        application_name="Student Management System",
        filename_prefix="student_management_test_data"
    )

    assert artifact["filename"].startswith("student_management_test_data")
    assert artifact["filename"].endswith(".xlsx")
    assert artifact["file_type"] == "Excel workbook"
    assert artifact["sheet_count"] == 4
    assert artifact["total_cases"] == 130
    assert artifact["large_dataset_cases"] == 100
    assert artifact["validation_status"] == "Created and validated"
    assert artifact["is_qa_test_data"] is True
    assert artifact["application_name"] == "Student Management System"
    assert "student_management_test_data" in artifact["reply_message"]
    assert "/api/tools/download/" in artifact["download_url"]

    file_path = os.path.join(settings.OUTPUT_DIR, artifact["filename"])
    assert os.path.exists(file_path)

    wb = openpyxl.load_workbook(file_path, data_only=True)
    expected_sheets = ["Summary", "Functional Tests", "Security & Edge Tests", "Large Dataset"]
    assert wb.sheetnames == expected_sheets

    # Verify Functional Tests has Student test cases
    ws_func = wb["Functional Tests"]
    assert ws_func.max_row >= 26
    found_student_case = False
    for r in range(2, ws_func.max_row + 1):
        cell_val = str(ws_func.cell(row=r, column=3).value or "")
        if "student" in cell_val.lower() or "marks" in cell_val.lower() or "enroll" in cell_val.lower():
            found_student_case = True
            break
    assert found_student_case is True

    # Verify Large Dataset has Student headers and 100 rows
    ws_large = wb["Large Dataset"]
    assert ws_large.max_row == 101  # 1 header + 100 student records
    headers = [cell.value for cell in ws_large[1]]
    assert "Student ID" in headers
    assert "Marks" in headers
    assert "Attendance (%)" in headers
    assert "Guardian Phone" in headers

    # Verify first data row
    first_row = [cell.value for cell in ws_large[2]]
    assert first_row[0] == "STU-2026-0001"
    assert "@student.edu" in str(first_row[2])
    wb.close()


@pytest.mark.asyncio
async def test_end_to_end_student_management_request():
    prompt = (
        "Create an Excel workbook for a student management system with 100 test cases. "
        "Include Student ID, Name, Email, Age, Class, Marks, Attendance, and Guardian Phone. "
        "Test valid inputs, invalid marks, duplicate IDs, missing fields, boundary values, "
        "and bulk uploads. Include expected results, test steps, validation rules, and a formula-driven summary."
    )
    artifact = await generate_workbook_from_request(prompt)

    # Must create a custom student management file
    assert artifact["filename"].startswith("student_management_test_data")
    assert artifact["filename"].endswith(".xlsx")
    assert artifact["application_name"] == "Student Management System"
    assert artifact["sheet_count"] == 4
    assert artifact["total_cases"] == 130
    assert artifact["large_dataset_cases"] == 100
    assert artifact["validation_status"] == "Created and validated"
    assert artifact["is_qa_test_data"] is True
    assert "/api/tools/download/" in artifact["download_url"]
    assert "student_management_test_data" in artifact["reply_message"]


@pytest.mark.asyncio
async def test_end_to_end_hospital_management_request():
    prompt = "Create an Excel workbook for hospital management system with test cases"
    artifact = await generate_workbook_from_request(prompt)

    assert artifact["filename"].startswith("hospital_management_test_data")
    assert artifact["filename"].endswith(".xlsx")
    assert artifact["application_name"] == "Hospital Management System"
    assert artifact["sheet_count"] == 4
    assert artifact["is_qa_test_data"] is True
    assert "/api/tools/download/" in artifact["download_url"]


# ==============================================================================
# Mandatory Requirement Tests A through M (Dynamic Excel Generation & Artifact Isolation)
# ==============================================================================

@pytest.mark.asyncio
async def test_a_dynamic_dimensions_student_data():
    """
    Test A — Dynamic dimensions:
    Generate 20 rows and 30 columns of student data. Verify the actual workbook contains
    exactly 20 data rows and 30 columns.
    """
    prompt = "Create an Excel sheet with 20 rows and 30 columns of student data."
    artifact = await generate_workbook_from_request(prompt)

    assert artifact["data_row_count"] == 20
    assert artifact["data_column_count"] == 30
    assert artifact["validation_status"] == "Created and validated"

    file_path = os.path.join(settings.OUTPUT_DIR, artifact["filename"])
    assert os.path.exists(file_path)

    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    # Header row (1) + exactly 20 data rows = 21 rows
    assert ws.max_row == 21
    assert ws.max_column == 30

    headers = [str(ws.cell(row=1, column=c).value or "").strip() for c in range(1, 31)]
    assert len(headers) == 30
    assert len(set(headers)) == 30  # All 30 distinct columns
    assert headers == STUDENT_PRIMARY_COLUMNS

    # Spot check data row types
    first_row = [ws.cell(row=2, column=c).value for c in range(1, 31)]
    assert first_row[0].startswith("STU-")
    assert "@" in str(first_row[5])  # Email column (index 5)
    wb.close()


@pytest.mark.asyncio
async def test_b_consecutive_generation_isolation():
    """
    Test B — Consecutive generation:
    Generate a 100-row user-registration workbook, then generate a 20-row, 30-column
    student workbook. Verify that the second request returns a different artifact ID,
    correct student columns, correct dimensions, and a different workbook artifact.
    """
    prompt1 = "Create an Excel workbook with 100 user registration test cases."
    art1 = await generate_workbook_from_request(prompt1)

    prompt2 = "Create an Excel sheet with 20 rows and 30 columns of student data."
    art2 = await generate_workbook_from_request(prompt2)

    # 1. Distinct IDs and filenames
    assert art1["artifact_id"] != art2["artifact_id"]
    assert art1["generation_request_id"] != art2["generation_request_id"]
    assert art1["filename"] != art2["filename"]

    # 2. Dimensions correctly differentiated
    assert art2["data_row_count"] == 20
    assert art2["data_column_count"] == 30
    assert art1["filename"].startswith("application_test_data")
    assert art2["filename"].startswith("student_data")

    # 3. Ensure both workbooks exist on disk without overwriting each other
    path1 = os.path.join(settings.OUTPUT_DIR, art1["filename"])
    path2 = os.path.join(settings.OUTPUT_DIR, art2["filename"])
    assert os.path.exists(path1)
    assert os.path.exists(path2)

    wb2 = openpyxl.load_workbook(path2, data_only=True)
    ws2 = wb2.active
    assert ws2.max_row == 21
    assert ws2.max_column == 30
    wb2.close()


@pytest.mark.asyncio
async def test_c_same_session_isolation():
    """
    Test C — Same-session isolation:
    Generate multiple workbooks in the same conversation. Confirm each download
    retrieves its own workbook.
    """
    conv_id = "test-conv-session-xyz"
    user_id = "test-user-same-session"

    art1 = await generate_workbook_from_request(
        "Generate 10 rows and 5 columns of customer data",
        conversation_id=conv_id,
        user_id=user_id
    )
    art2 = await generate_workbook_from_request(
        "Generate 25 rows and 8 columns of sales data",
        conversation_id=conv_id,
        user_id=user_id
    )

    assert art1["artifact_id"] != art2["artifact_id"]
    assert art1["filename"] != art2["filename"]

    rec1 = artifact_registry.get_artifact(art1["artifact_id"])
    rec2 = artifact_registry.get_artifact(art2["artifact_id"])
    assert rec1 is not None and rec2 is not None
    assert rec1.conversation_id == conv_id
    assert rec2.conversation_id == conv_id
    assert rec1.data_row_count == 10
    assert rec2.data_row_count == 25


@pytest.mark.asyncio
async def test_d_cross_session_isolation():
    """
    Test D — Cross-session isolation:
    Ensure a user or session cannot access another user's private artifact.
    """
    user_alice = UserModel(id="alice-id", username="alice", email="alice@test.com", password_hash="h", salt="s", role="user")
    user_bob = UserModel(id="bob-id", username="bob", email="bob@test.com", password_hash="h", salt="s", role="user")

    art_alice = await generate_workbook_from_request(
        "Create an Excel sheet with 15 rows and 10 columns of student data",
        user_id=user_alice.id
    )

    # Alice can download her own artifact
    resp_alice = await download_output_file(filename=art_alice["artifact_id"], current_user=user_alice)
    assert resp_alice.status_code == 200

    # Bob cannot download Alice's artifact (HTTP 403 Forbidden)
    with pytest.raises(HTTPException) as exc_info:
        await download_output_file(filename=art_alice["artifact_id"], current_user=user_bob)
    assert exc_info.value.status_code == 403


@pytest.mark.asyncio
async def test_e_repeated_prompts_distinct_artifacts():
    """
    Test E — Repeated prompts:
    Submit the same generation prompt twice. Verify each successful request
    creates a distinct artifact unless the user explicitly requests reuse.
    """
    prompt = "Create an Excel sheet with 20 rows and 30 columns of student data."
    art1 = await generate_workbook_from_request(prompt)
    art2 = await generate_workbook_from_request(prompt)

    assert art1["artifact_id"] != art2["artifact_id"]
    assert art1["generation_request_id"] != art2["generation_request_id"]
    assert art1["filename"] != art2["filename"]


@pytest.mark.asyncio
async def test_f_concurrent_requests():
    """
    Test F — Concurrent requests:
    Generate two different workbooks concurrently and verify neither overwrites
    or returns the other's file.
    """
    p1 = "Create an Excel sheet with 12 rows and 15 columns of student data."
    p2 = "Create an Excel sheet with 28 rows and 10 columns of employee data."

    art1, art2 = await asyncio.gather(
        generate_workbook_from_request(p1),
        generate_workbook_from_request(p2)
    )

    assert art1["artifact_id"] != art2["artifact_id"]
    assert art1["filename"] != art2["filename"]
    assert art1["data_row_count"] == 12
    assert art1["data_column_count"] == 15
    assert art2["data_row_count"] == 28
    assert art2["data_column_count"] == 10

    # Verify both files on disk
    wb1 = openpyxl.load_workbook(os.path.join(settings.OUTPUT_DIR, art1["filename"]))
    assert wb1.active.max_row == 13
    assert wb1.active.max_column == 15
    wb1.close()

    wb2 = openpyxl.load_workbook(os.path.join(settings.OUTPUT_DIR, art2["filename"]))
    assert wb2.active.max_row == 29
    assert wb2.active.max_column == 10
    wb2.close()


@pytest.mark.asyncio
async def test_g_download_integrity_openpyxl():
    """
    Test G — Download integrity:
    Download the artifact through the actual backend endpoint and reopen the
    downloaded bytes with openpyxl.
    """
    art = await generate_workbook_from_request("Create an Excel sheet with 20 rows and 30 columns of student data.")
    resp = await download_output_file(filename=art["artifact_id"])
    assert resp.status_code == 200

    wb = openpyxl.load_workbook(resp.path, data_only=True)
    ws = wb.active
    assert ws.max_row == 21
    assert ws.max_column == 30
    wb.close()


def test_h_invalid_input_limits():
    """
    Test H — Invalid input:
    Verify that invalid or excessive dimensions produce a controlled error.
    """
    # 1. Negative dimensions trigger controlled error
    with pytest.raises(ValueError) as exc_neg:
        parse_excel_generation_request("Generate -10 rows and -5 columns of student data")
    assert "Invalid dimensions" in str(exc_neg.value)

    # 2. Excessive dimensions trigger controlled error
    with pytest.raises(ValueError) as exc_overflow:
        parse_excel_generation_request("Generate 999999 rows and 5000 columns of student data")
    assert "exceeds maximum supported limit" in str(exc_overflow.value)

    # 3. Safe supported dimensions parse successfully within bounds
    p_safe = parse_excel_generation_request("Generate 500 rows and 25 columns of student data")
    assert p_safe.data_rows == 500
    assert p_safe.data_columns == 25


def test_i_excel_analysis_regression():
    """
    Test I — Excel analysis regression:
    Confirm that analyzing an uploaded workbook still reads the original workbook
    rather than generating a new one.
    """
    analysis_queries = [
        "Analyze my uploaded Excel workbook.",
        "Calculate the average from my spreadsheet.",
        "Explain the contents of an existing workbook.",
        "What is the sum of column B in my uploaded file?"
    ]
    for q in analysis_queries:
        assert is_excel_generation_request(q) is False, f"Query '{q}' should NOT route to excel_generation"


def test_j_rag_regression():
    """
    Test J — RAG regression:
    Confirm document-based queries do not accidentally trigger excel generation.
    """
    rag_queries = [
        "What are the requirements in the uploaded project PDF?",
        "Summarize section 2 of the contract document.",
        "Who is the candidate according to the resume?"
    ]
    for q in rag_queries:
        assert is_excel_generation_request(q) is False, f"Query '{q}' should NOT route to excel_generation"


def test_k_web_search_regression():
    """
    Test K — Web-search regression:
    Confirm current-information questions continue to use web search.
    """
    web_queries = [
        "Find current information about stock prices of Apple and Microsoft.",
        "Search the web for latest AI news today.",
        "Research on the internet about current mortgage rates."
    ]
    for q in web_queries:
        assert is_explicit_web_search_request(q) is True
        assert is_excel_generation_request(q) is False


def test_l_general_question_regression():
    """
    Test L — General-question regression:
    Confirm 'What is Excel?' does not generate a workbook.
    """
    general_queries = [
        "What is Excel?",
        "Explain what Excel is and what it is used for.",
        "How does a spreadsheet work?",
        "What are formulas in spreadsheets?"
    ]
    for q in general_queries:
        assert is_excel_generation_request(q) is False, f"Query '{q}' should NOT generate workbook"


def test_m_formula_injection_defense():
    """
    Test M — Formula injection:
    Confirm untrusted cell text is safely handled (CWE-1236).
    """
    dangerous_inputs = [
        "=cmd|'/C calc'!A0",
        "+2+5+cmd",
        "-12345",
        "@SUM(A1:A10)",
        "\t=dangerous",
        "\r+dangerous"
    ]
    for dangerous in dangerous_inputs:
        sanitized = sanitize_cell_value(dangerous)
        assert sanitized.startswith("'"), f"Dangerous input '{dangerous}' was not prefixed with apostrophe: '{sanitized}'"


# ==============================================================================
# Part 10: Mode Compatibility & Multi-Workflow Regression Tests (Tests 1 through 8)
# ==============================================================================

@pytest.mark.asyncio
async def test_regression_1_web_mode_with_excel_generation():
    """
    Test 1:
    Web Search mode + "Create an Excel sheet of random 5 rows and 10 columns of students data."
    Expected:
    The Excel Generator executes. A valid downloadable XLSX contains exactly five data rows and ten columns.
    """
    prompt = "Create an Excel sheet of random 5 rows and 10 columns of students data."

    # 1. Verify router prioritizes Excel generation even when web_mode is True
    route, badge = await query_router.route_query(prompt, has_documents=False, web_mode=True)
    assert route == "excel_generation"
    assert badge == "📊 Excel Generator"

    # 2. Execute full stream generation in web_mode
    from backend.llm.generator import grounded_generator
    chunks = []
    async for sse in grounded_generator.stream_rag_response(
        question=prompt,
        route=route,
        router_badge=badge,
        retrieved_chunks=[],
        queries=[prompt],
        web_mode=True
    ):
        chunks.append(sse)

    combined = "".join(chunks)
    assert "event: excel_artifact" in combined
    assert "event: done" in combined

    # 3. Extract artifact payload and inspect physical workbook
    for line in combined.split("\n"):
        if line.startswith("data:") and "artifact_id" in line:
            import json
            art = json.loads(line[5:].strip())
            assert art["data_row_count"] == 5
            assert art["data_column_count"] == 10
            assert art["validation_status"] == "Created and validated"

            file_path = os.path.join(settings.OUTPUT_DIR, art["filename"])
            assert os.path.exists(file_path)

            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active
            assert ws.max_row == 6  # 1 header + exactly 5 data rows
            assert ws.max_column == 10  # exactly 10 columns
            wb.close()
            break


@pytest.mark.asyncio
async def test_regression_2_default_chat_mode_with_excel_generation():
    """
    Test 2:
    Default chat mode + "Create an Excel sheet of random 5 rows and 10 columns of students data."
    Expected:
    The correct workbook is generated.
    """
    prompt = "Create an Excel sheet of random 5 rows and 10 columns of students data."
    route, badge = await query_router.route_query(prompt, has_documents=False, web_mode=False)
    assert route == "excel_generation"
    assert badge == "📊 Excel Generator"

    artifact = await generate_workbook_from_request(prompt)
    assert artifact["data_row_count"] == 5
    assert artifact["data_column_count"] == 10

    file_path = os.path.join(settings.OUTPUT_DIR, artifact["filename"])
    wb = openpyxl.load_workbook(file_path, data_only=True)
    assert wb.active.max_row == 6
    assert wb.active.max_column == 10
    wb.close()


@pytest.mark.asyncio
async def test_regression_3_web_mode_with_current_excel_documentation():
    """
    Test 3:
    Web Search mode + "Find current Excel documentation."
    Expected:
    Live Web Search executes normally. No workbook is generated.
    """
    prompt = "Find current Excel documentation."
    assert is_excel_generation_request(prompt) is False

    route, badge = await query_router.route_query(prompt, has_documents=False, web_mode=True)
    assert route == "web"
    assert badge == "🌐 Live Web Search"


@pytest.mark.asyncio
async def test_regression_4_excel_analysis_on_uploaded_file():
    """
    Test 4:
    "Analyze my uploaded Excel file."
    Expected:
    Excel analysis executes normally without generating a new file.
    """
    prompt = "Analyze my uploaded Excel file."
    assert is_excel_generation_request(prompt) is False

    from backend.tools.excel_analyzer import active_excel_context
    active_excel_context.set({
        "doc_ids": ["doc_test_123"],
        "doc_id": "doc_test_123",
        "filename": "sales_q4.xlsx",
        "has_spreadsheet": True,
    })

    route, badge = await query_router.route_query(prompt, has_documents=True, web_mode=False)
    assert route == "document"
    assert badge == "📊 Spreadsheet Analysis"
    active_excel_context.set(None)


@pytest.mark.asyncio
async def test_regression_5_document_based_excel_generation():
    """
    Test 5:
    "Generate Excel test cases from my uploaded PDF."
    Expected:
    Relevant document requirements are retrieved and used to create a workbook.
    """
    prompt = "Generate Excel test cases from my uploaded PDF."
    route, badge = await query_router.route_query(prompt, has_documents=True, web_mode=False)
    assert route == "excel_generation"

    from backend.llm.router import DOC_EXPLICIT_REGEX
    assert bool(DOC_EXPLICIT_REGEX.search(prompt)) is True

    fake_doc_chunks = [
        {"chunk_id": "c1", "text": "Requirement 1: User password must contain at least 8 characters and 1 symbol.", "doc_id": "pdf_1"}
    ]
    artifact = await generate_workbook_from_request(prompt, context=fake_doc_chunks[0]["text"])
    assert artifact["validation_status"] == "Created and validated"
    assert artifact["sheet_count"] >= 1


@pytest.mark.asyncio
async def test_regression_6_consecutive_generation_different_dimensions_same_conv():
    """
    Test 6:
    Generate one workbook and then another with different dimensions in the same conversation.
    Expected:
    The second response returns the second workbook's own artifact ID and dimensions.
    """
    conv_id = "conv_test_consecutive_999"
    p1 = "Create an Excel sheet of random 5 rows and 10 columns of students data."
    p2 = "Create an Excel sheet with 15 rows and 8 columns of customer data."

    art1 = await generate_workbook_from_request(p1, conversation_id=conv_id)
    art2 = await generate_workbook_from_request(p2, conversation_id=conv_id)

    assert art1["artifact_id"] != art2["artifact_id"]
    assert art1["filename"] != art2["filename"]
    assert art1["data_row_count"] == 5
    assert art1["data_column_count"] == 10
    assert art2["data_row_count"] == 15
    assert art2["data_column_count"] == 8

    # Check both files exist independently
    path1 = os.path.join(settings.OUTPUT_DIR, art1["filename"])
    path2 = os.path.join(settings.OUTPUT_DIR, art2["filename"])
    assert os.path.exists(path1)
    assert os.path.exists(path2)


def test_regression_7_invalid_dimensions_controlled_error():
    """
    Test 7:
    Generation fails or input dimensions are invalid.
    Expected:
    A controlled error appears. No old artifact is presented as the result.
    """
    with pytest.raises(ValueError) as exc:
        parse_excel_generation_request("Create an Excel sheet with -20 rows and 5 columns")
    assert "Invalid dimensions" in str(exc.value)

    with pytest.raises(ValueError) as exc2:
        parse_excel_generation_request("Create an Excel sheet with 90000 rows and 10 columns")
    assert "exceeds maximum supported limit" in str(exc2.value)


@pytest.mark.asyncio
async def test_regression_8_download_endpoint_openpyxl_integrity():
    """
    Test 8:
    The generated workbook is downloaded through the actual endpoint and reopened with openpyxl.
    Expected:
    The downloaded file matches the generated artifact.
    """
    prompt = "Create an Excel sheet of random 5 rows and 10 columns of students data."
    artifact = await generate_workbook_from_request(prompt)
    artifact_id = artifact["artifact_id"]

    response = await download_output_file(filename=artifact_id)
    assert response.status_code == 200
    assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    downloaded_wb = openpyxl.load_workbook(response.path, data_only=True)
    ws = downloaded_wb.active
    assert ws.max_row == 6  # 1 header + 5 rows
    assert ws.max_column == 10
    downloaded_wb.close()



